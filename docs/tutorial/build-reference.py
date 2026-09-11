from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path('C:/Users/Public/drawer-performance-344');src=root/'docs/tutorial/서랍마을-전체기능-튜토리얼기획-348.md';out=Path('C:/Users/김세은/OneDrive/문서/평행도시 게임 만들기/앱 전달');out.mkdir(exist_ok=True)
d=Document();section=d.sections[0];section.page_width=Inches(8.27);section.page_height=Inches(11.69);section.top_margin=Inches(.65);section.bottom_margin=Inches(.65);section.left_margin=section.right_margin=Inches(.6)
for n,size in [('Normal',10),('Title',23),('Heading 1',16),('Heading 2',12)]:
 st=d.styles[n];st.font.name='맑은 고딕';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕');st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=1.05
for n in ['Heading 1','Heading 2']:d.styles[n].paragraph_format.space_before=Pt(7)
d.styles['Normal'].paragraph_format.widow_control=True
header=section.header.paragraphs[0];header.text='서랍마을  |  튜토리얼 기획 참고서';header.style='Normal';header.runs[0].font.size=Pt(8)
footer=section.footer.paragraphs[0];footer.alignment=2;footer.add_run('348 개발 버전  ·  ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
for line in src.read_text(encoding='utf-8-sig').splitlines():
 line=line.strip()
 if not line:continue
 if line=='<!-- page -->':d.add_page_break();continue
 if line.startswith('# '):d.add_paragraph(line[2:],style='Title')
 elif line.startswith('## '):d.add_paragraph(line[3:],style='Heading 1')
 elif line.startswith('### '):d.add_paragraph(line[4:],style='Heading 2')
 elif line.startswith('- '):
  p=d.add_paragraph('• '+line[2:]);p.paragraph_format.left_indent=Inches(.1);p.paragraph_format.first_line_indent=Inches(-.1)
 else:d.add_paragraph(line)
dest=out/'서랍마을-전체기능-튜토리얼기획-348.docx';d.save(dest)
(out/src.name).write_text(src.read_text(encoding='utf-8-sig').replace('<!-- page -->\n',''),encoding='utf-8')
print(dest)
